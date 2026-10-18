from __future__ import annotations

"""Editable PowerPoint rendering for deck deliverables.

Uses python-pptx to produce a native .pptx made of real text boxes, bullet
lists and tables — no flattened images — so reviewers can open the deck in
PowerPoint/Keynote/Google Slides and edit every word. Layout follows the
Steer design system colors (pipeline/templates/design-system.json).
"""

from pathlib import Path

import re

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Inches, Pt

from app.services.rendering.fonts import design_tokens

TOKENS = design_tokens()
C = TOKENS["colors"]

SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)

FONT_HEAD = "Georgia"
FONT_BODY = "Calibri"


def _rgb(hex_color: str) -> RGBColor:
    h = hex_color.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


def _blank_slide(prs: Presentation, bg: str):
    layout = prs.slide_layouts[6]  # blank
    slide = prs.slides.add_slide(layout)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = _rgb(bg)
    return slide


def _textbox(slide, left, top, width, height):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    return box, tf


def _set_para(p, text, *, size, color, bold=False, font=FONT_BODY, space_after=6, align=None):
    p.text = text
    if align is not None:
        p.alignment = align
    for run in p.runs:
        run.font.size = Pt(size)
        run.font.color.rgb = _rgb(color)
        run.font.bold = bold
        run.font.name = font
    p.space_after = Pt(space_after)


def _brandmark(slide, dark: bool = False):
    _, tf = _textbox(slide, SLIDE_W - Inches(1.6), SLIDE_H - Inches(0.55), Inches(1.3), Inches(0.4))
    _set_para(tf.paragraphs[0], TOKENS["brandmark"], size=14, color=C["cream"] if dark else C["clay"], bold=True, font=FONT_HEAD)


def _footer_note(slide, text: str, dark: bool = False):
    _, tf = _textbox(slide, Inches(0.6), SLIDE_H - Inches(0.55), SLIDE_W - Inches(2.4), Inches(0.4))
    _set_para(tf.paragraphs[0], text[:180], size=10, color=C["blush"] if dark else C["clay"])


def _clean_bullet(text: str) -> str:
    """Normalise a bullet so a truncated fragment ("O", "Ob", "Bas") never reaches a slide."""
    t = " ".join(str(text or "").split()).strip()
    t = re.sub(r"^[•·\-\u2022\u00b7]+", "", t).strip()
    if not t:
        return ""
    if len(t) < 5:
        return ""
    letters = sum(1 for ch in t if ch.isalpha())
    if letters < 5:
        return ""
    if letters == len(t) and not re.search(r"[a-z]", t):
        return ""
    return t


def _bullets(items: list[str], limit: int = 8) -> list[str]:
    out: list[str] = []
    seen: set[str] = set()
    for item in items:
        b = _clean_bullet(item)
        if not b or b.lower() in seen:
            continue
        seen.add(b.lower())
        out.append(b)
        if len(out) >= limit:
            break
    return out


def render_deck(payload: dict, out_dir: Path) -> list[dict]:
    """Render the deck payload to a single native .pptx. Returns asset specs."""
    out_dir.mkdir(parents=True, exist_ok=True)
    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H

    deck_title = (payload.get("deck_title") or "Research briefing").strip()
    subtitle = (payload.get("subtitle") or "").strip()
    disclaimer = (payload.get("disclaimer") or "").strip()

    slides = payload.get("slides") or []
    for i, sl in enumerate(slides):
        kind = (sl.get("kind") or "content").strip()
        title = (sl.get("title") or f"Slide {i + 1}").strip()
        bullets = _bullets(sl.get("bullets") or [])
        table = sl.get("table")
        claim_ids = sl.get("claim_ids") or []
        source_notes = sl.get("source_notes") or []

        if kind == "title":
            _render_title_slide(prs, deck_title, subtitle, bullets, source_notes)
            continue

        is_closing = kind in {"limitations", "takeaways", "references"}
        bg = C["cream"]
        fg = C["ink"]
        accent = C["rust"]
        if is_closing and kind == "references":
            bg, fg, accent = C["maroon"], C["cream"], C["blush"]

        slide = _blank_slide(prs, bg)
        # accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.6), Inches(0.55), Inches(0.18), Inches(0.7))
        bar.fill.solid()
        bar.fill.fore_color.rgb = _rgb(accent)
        bar.line.fill.background()

        # slide number chip
        _, num_tf = _textbox(slide, SLIDE_W - Inches(1.1), Inches(0.5), Inches(0.8), Inches(0.35))
        _set_para(num_tf.paragraphs[0], f"{i + 1}/{len(slides)}", size=11, color=accent, bold=True, align=PP_ALIGN.RIGHT)

        # title
        _, t_tf = _textbox(slide, Inches(1.0), Inches(0.45), SLIDE_W - Inches(2.2), Inches(0.9))
        _set_para(t_tf.paragraphs[0], title, size=30, color=fg, bold=True, font=FONT_HEAD)

        # rule
        rule = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.0), Inches(1.35), Inches(4.5), Emu(38100))
        rule.fill.solid()
        rule.fill.fore_color.rgb = _rgb(accent)
        rule.line.fill.background()

        y = Inches(1.6)
        content_h = SLIDE_H - Inches(2.4) - y

        if table and table.get("headers") and table.get("rows"):
            rows = len(table["rows"]) + 1
            cols = len(table["headers"])
            shape = slide.shapes.add_table(
                rows, cols, Inches(1.0), y, SLIDE_W - Inches(2.0), Inches(0.4 * rows)
            )
            tbl = shape.table
            for c, h in enumerate(table["headers"]):
                cell = tbl.cell(0, c)
                cell.text = str(h)
                for p in cell.text_frame.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(13)
                        run.font.bold = True
                        run.font.color.rgb = _rgb(C["cream"])
                        run.font.name = FONT_BODY
                cell.fill.solid()
                cell.fill.fore_color.rgb = _rgb(C["maroon"])
            for r, row in enumerate(table["rows"], start=1):
                for c, val in enumerate(row[:cols]):
                    cell = tbl.cell(r, c)
                    cell.text = str(val)
                    for p in cell.text_frame.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(12)
                            run.font.color.rgb = _rgb(fg)
                            run.font.name = FONT_BODY
            y = y + Inches(0.4 * rows) + Inches(0.2)
            bullets = bullets[: max(0, 6 - rows)]  # keep the slide readable

        if bullets:
            _, b_tf = _textbox(slide, Inches(1.0), y, SLIDE_W - Inches(2.0), content_h)
            first = True
            for b in bullets[:8]:
                p = b_tf.paragraphs[0] if first else b_tf.add_paragraph()
                first = False
                _set_para(p, "• " + b, size=16, color=fg, space_after=8)

        foot = ""
        if claim_ids:
            foot += "Evidence: " + ", ".join(claim_ids)
        if source_notes:
            foot += (" · " if foot else "") + "; ".join(str(s) for s in source_notes[:3])
        if not foot and disclaimer and kind == "limitations":
            foot = disclaimer
        if foot:
            _footer_note(slide, foot, dark=(bg == C["maroon"]))
        _brandmark(slide, dark=(bg == C["maroon"]))

    if not slides:
        slide = _blank_slide(prs, C["cream"])
        _, tf = _textbox(slide, Inches(1.0), Inches(3.0), SLIDE_W - Inches(2.0), Inches(1.0))
        _set_para(tf.paragraphs[0], "This deck has no slides yet.", size=24, color=C["ink"], font=FONT_HEAD)

    out_path = out_dir / f"{_safe_name(deck_title)}.pptx"
    prs.save(str(out_path))
    return [{
        "kind": "deck_pptx",
        "path": str(out_path),
        "filename": out_path.name,
        "mime": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        "label": "Editable PowerPoint deck (.pptx)",
    }]


def _render_title_slide(prs: Presentation, deck_title: str, subtitle: str, bullets: list[str], source_notes: list[str]) -> None:
    slide = _blank_slide(prs, C["maroon"])
    _, tf = _textbox(slide, Inches(1.0), Inches(2.2), SLIDE_W - Inches(2.0), Inches(1.6))
    _set_para(tf.paragraphs[0], deck_title, size=40, color=C["cream"], bold=True, font=FONT_HEAD)
    if subtitle:
        p = tf.add_paragraph()
        _set_para(p, subtitle, size=18, color=C["blush"], font=FONT_BODY, space_after=10)
    if bullets:
        _, b_tf = _textbox(slide, Inches(1.0), Inches(4.2), SLIDE_W - Inches(2.0), Inches(2.2))
        first = True
        for b in bullets[:5]:
            p = b_tf.paragraphs[0] if first else b_tf.add_paragraph()
            first = False
            _set_para(p, "• " + b, size=13, color=C["blush"], space_after=5)
    _brandmark(slide, dark=True)


def _safe_name(title: str) -> str:
    keep = [c if (c.isalnum() or c in "-_ ") else "_" for c in title]
    name = "_".join("".join(keep).split())[:60].strip("_")
    return name or "steer_deck"


def pptx_available() -> bool:
    try:
        import pptx  # noqa: F401
        return True
    except Exception:
        return False
